"""CSV and Excel export."""
from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import LEAD_FIELDS, LEAD_LABELS, NOT_FOUND

_PHONE_LIKE = re.compile(r"^[+\-]?[0-9 ().\-]+$")


def _safe_cell(value):
    """Neutralise spreadsheet formula injection (=, @, and +/- unless it is just a phone number)."""
    if not isinstance(value, str) or not value:
        return value
    if value[0] in ("=", "@") or (value[0] in ("+", "-") and not _PHONE_LIKE.match(value)):
        return "'" + value
    return value


def leads_to_dataframe(leads: list[dict]) -> pd.DataFrame:
    rows = []
    for lead in leads:
        rows.append({label: _safe_cell(lead.get(key) or NOT_FOUND) for key, label in LEAD_FIELDS})
    return pd.DataFrame(rows, columns=LEAD_LABELS)


def export_csv(leads: list[dict]) -> bytes:
    # utf-8 with BOM so Excel on Windows opens accented characters correctly
    return leads_to_dataframe(leads).to_csv(index=False).encode("utf-8-sig")


def export_xlsx(leads: list[dict], meta: dict | None = None) -> bytes:
    df = leads_to_dataframe(leads)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Leads", index=False)
        ws = writer.sheets["Leads"]

        header_fill = PatternFill("solid", fgColor="1F3A5F")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for idx, column in enumerate(df.columns, start=1):
            longest = max([len(str(column))] + [len(str(v)) for v in df[column].head(500)])
            ws.column_dimensions[get_column_letter(idx)].width = max(14, min(longest + 2, 60))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

        notes = writer.book.create_sheet("About this export")
        info = [
            ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Search prompt", (meta or {}).get("prompt", "")),
            ("Leads", len(df)),
            ("Not Found", "Fields the sources did not publish are marked 'Not Found'; nothing was guessed."),
            ("Verification", "Statuses describe where a record came from. Phone numbers and emails are NOT "
                             "tested; do not treat them as validated."),
            ("Compliance", "Public business details only. Check anti-spam / data-protection rules (e.g. CASL, "
                           "CAN-SPAM, GDPR) before contacting anyone."),
        ]
        for r, (k, v) in enumerate(info, start=1):
            notes.cell(row=r, column=1, value=k).font = Font(bold=True)
            notes.cell(row=r, column=2, value=_safe_cell(v) if isinstance(v, str) else v)
        notes.column_dimensions["A"].width = 16
        notes.column_dimensions["B"].width = 100
    return buffer.getvalue()
