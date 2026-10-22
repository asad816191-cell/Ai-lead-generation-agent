import csv
import io
import unittest

from openpyxl import load_workbook

from services.export_service import export_csv, export_xlsx
from services.models import LEAD_LABELS, NOT_FOUND, new_lead

LEADS = [
    new_lead(company_name="Acme Realty", phone="+1 416-555-0100", email="info@acme.ca",
             website="https://acme.ca", address="1 King St W, Toronto, ON", city="Toronto",
             region="Ontario", country="Canada", source_url="https://acme.ca/contact",
             source_type="Official website", verification_status="Official website visited"),
    new_lead(company_name="=HYPERLINK(\"http://evil\")", phone="+1 647-555-0101"),
]


class ExportTests(unittest.TestCase):
    def test_csv_has_all_fields_and_not_found(self):
        data = export_csv(LEADS)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))  # BOM for Excel
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(rows[0], LEAD_LABELS)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1][0], "Acme Realty")
        self.assertEqual(rows[2][LEAD_LABELS.index("Public Business Email")], NOT_FOUND)

    def test_csv_formula_injection_neutralised_but_phone_kept(self):
        rows = list(csv.reader(io.StringIO(export_csv(LEADS).decode("utf-8-sig"))))
        self.assertTrue(rows[2][0].startswith("'="))
        self.assertEqual(rows[2][LEAD_LABELS.index("Public Business Phone")], "+1 647-555-0101")

    def test_xlsx_formatting(self):
        wb = load_workbook(io.BytesIO(export_xlsx(LEADS, {"prompt": "Find 2 realtors in Toronto"})))
        ws = wb["Leads"]
        self.assertEqual([c.value for c in ws[1]], LEAD_LABELS)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertIsNotNone(ws.auto_filter.ref)
        self.assertEqual(ws.max_row, 3)
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 14)
        self.assertGreater(ws.column_dimensions["G"].width, ws.column_dimensions["B"].width - 1)
        self.assertIn("About this export", wb.sheetnames)

    def test_empty_export(self):
        wb = load_workbook(io.BytesIO(export_xlsx([])))
        self.assertEqual(wb["Leads"].max_row, 1)


if __name__ == "__main__":
    unittest.main()
